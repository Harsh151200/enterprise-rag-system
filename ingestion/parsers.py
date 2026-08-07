import io
import os
from typing import Union
from bs4 import BeautifulSoup
import pypdf
import docx
import openpyxl
import pptx

from ingestion.base import BaseParser, ExtractedDocument

class TXTParser(BaseParser):
    """Handles raw plain text and standard markdown files."""
    def parse(self, raw_content: Union[bytes, str], file_name: str) -> ExtractedDocument:
        if isinstance(raw_content, bytes):
            text = raw_content.decode("utf-8", errors="ignore")
        else:
            text = raw_content
            
        metadata = {"source_file": file_name, "format": "txt"}
        return ExtractedDocument(raw_text=text.strip(), metadata=metadata)


class HTMLParser(BaseParser):
    """Handles text extraction from web pages and HTML documentation nodes."""
    def parse(self, raw_content: Union[bytes, str], file_name: str) -> ExtractedDocument:
        if isinstance(raw_content, bytes):
            raw_content = raw_content.decode("utf-8", errors="ignore")
            
        soup = BeautifulSoup(raw_content, "lxml")
        
        # Prioritize primary content boundaries if common structural tags exist
        main_body = soup.find(["article", "main", "div.content"])
        if main_body:
            text_content = main_body.get_text(separator="\n")
        else:
            text_content = soup.get_text(separator="\n")
            
        clean_lines = [line.strip() for line in text_content.splitlines() if line.strip()]
        final_text = "\n".join(clean_lines)
        
        metadata = {"source_file": file_name, "format": "html"}
        return ExtractedDocument(raw_text=final_text, metadata=metadata)


class PDFParser(BaseParser):
    """Handles extraction from Portable Document Format (PDF) files."""
    def parse(self, raw_content: Union[bytes, str], file_name: str) -> ExtractedDocument:
        if isinstance(raw_content, str):
            raw_content = raw_content.encode("utf-8")
            
        pdf_file_obj = io.BytesIO(raw_content)
        reader = pypdf.PdfReader(pdf_file_obj)
        
        extracted_pages = []
        for page in reader.pages:
            text = page.extract_text()
            if text:
                extracted_pages.append(text)
                
        final_text = "\n\n".join(extracted_pages)
        metadata = {"source_file": file_name, "format": "pdf", "total_pages": len(reader.pages)}
        return ExtractedDocument(raw_text=final_text.strip(), metadata=metadata)


class DOCXParser(BaseParser):
    """Handles unzipping and parsing text components out of OpenXML Microsoft Word records."""
    def parse(self, raw_content: Union[bytes, str], file_name: str) -> ExtractedDocument:
        if isinstance(raw_content, str):
            raw_content = raw_content.encode("utf-8")
            
        file_stream = io.BytesIO(raw_content)
        doc_obj = docx.Document(file_stream)
        
        paragraphs_text = [p.text for p in doc_obj.paragraphs if p.text.strip()]
        final_text = "\n".join(paragraphs_text)
        
        metadata = {"source_file": file_name, "format": "docx"}
        return ExtractedDocument(raw_text=final_text.strip(), metadata=metadata)


class ExcelParser(BaseParser):
    """Handles reading OpenXML spreadsheets, flattening data row-by-row."""
    def parse(self, raw_content: Union[bytes, str], file_name: str) -> ExtractedDocument:
        if isinstance(raw_content, str):
            raw_content = raw_content.encode("utf-8")
            
        file_stream = io.BytesIO(raw_content)
        workbook = openpyxl.load_workbook(file_stream, data_only=True, read_only=True)
        
        row_strings = []
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            row_strings.append(f"--- Sheet: {sheet_name} ---")
            for row in sheet.iter_rows(values_only=True):
                # Filter out completely empty spreadsheet cells
                filtered_values = [str(cell_val).strip() for cell_val in row if cell_val is not None]
                if filtered_values:
                    row_strings.append(" | ".join(filtered_values))
                    
        final_text = "\n".join(row_strings)
        metadata = {"source_file": file_name, "format": "xlsx"}
        return ExtractedDocument(raw_text=final_text.strip(), metadata=metadata)


class PPTXParser(BaseParser):
    """Handles looping through visual presentation shapes to harvest presentation texts."""
    def parse(self, raw_content: Union[bytes, str], file_name: str) -> ExtractedDocument:
        if isinstance(raw_content, str):
            raw_content = raw_content.encode("utf-8")
            
        file_stream = io.BytesIO(raw_content)
        presentation = pptx.Presentation(file_stream)
        
        slide_texts = []
        for idx, slide in enumerate(presentation.slides):
            slide_texts.append(f"--- Slide {idx + 1} ---")
            for shape in slide.shapes:
                if hasattr(shape, "text_frame") and shape.text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        if paragraph.text.strip():
                            slide_texts.append(paragraph.text.strip())
                            
        final_text = "\n".join(slide_texts)
        metadata = {"source_file": file_name, "format": "pptx", "total_slides": len(presentation.slides)}
        return ExtractedDocument(raw_text=final_text.strip(), metadata=metadata)


class XMLAndCodeParser(BaseParser):
    """Handles reading structural files and code components while capturing software metadata."""
    def parse(self, raw_content: Union[bytes, str], file_name: str) -> ExtractedDocument:
        if isinstance(raw_content, bytes):
            text = raw_content.decode("utf-8", errors="ignore")
        else:
            text = raw_content
            
        file_extension = os.path.splitext(file_name)[1].lower().replace(".", "")
        if not file_extension:
            file_extension = "code"
            
        metadata = {"source_file": file_name, "format": "structured_code", "language": file_extension}
        return ExtractedDocument(raw_text=text.strip(), metadata=metadata)